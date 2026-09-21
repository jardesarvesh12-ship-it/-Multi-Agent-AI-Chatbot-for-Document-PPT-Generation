"""
generate_samples.py — Generate sample DOCX and PPTX templates for testing.
Run: python scripts/generate_samples.py
"""
from pathlib import Path
import sys

# Configure UTF-8 output encoding on Windows consoles
if sys.platform == "win32" and hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

# Add project root to path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from pptx import Presentation
from pptx.util import Inches as PptInches, Pt as PptPt
from pptx.dml.color import RGBColor as PptRGBColor


def create_sample_docx():
    """Create sample Company_Proposal.docx template."""
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)

    # Set default font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # Configure heading styles
    h1 = doc.styles["Heading 1"]
    h1.font.name = "Calibri"
    h1.font.size = Pt(16)
    h1.font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)
    h1.font.bold = True

    h2 = doc.styles["Heading 2"]
    h2.font.name = "Calibri"
    h2.font.size = Pt(14)
    h2.font.color.rgb = RGBColor(0x2C, 0x3E, 0x50)

    # ── Document Content ──
    title = doc.add_heading("Technology Innovation Proposal", level=0)
    for run in title.runs:
        run.font.size = Pt(24)
        run.font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)

    doc.add_paragraph("Confidential — Prepared by Acme Corporation")
    doc.add_paragraph("")

    # Executive Summary
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(
        "This proposal outlines our strategic approach to leveraging cutting-edge "
        "technology solutions for enterprise transformation. Our team of experts has "
        "conducted extensive research to develop a comprehensive roadmap that addresses "
        "the key challenges and opportunities in the current market landscape."
    )
    doc.add_paragraph(
        "The proposed initiative will drive measurable outcomes across operational "
        "efficiency, customer engagement, and revenue growth through the adoption of "
        "AI-powered platforms and modern cloud infrastructure."
    )

    # Background
    doc.add_heading("Background & Context", level=1)
    doc.add_paragraph(
        "The digital transformation landscape continues to evolve rapidly. "
        "Organizations that fail to adapt risk falling behind their competitors. "
        "Our analysis identifies three critical areas requiring immediate attention:"
    )
    bullets = [
        "Cloud-native architecture migration for scalability and resilience",
        "AI and machine learning integration for data-driven decision making",
        "Enhanced cybersecurity posture to protect enterprise assets",
        "Modernized customer experience platforms for improved engagement",
    ]
    for b in bullets:
        doc.add_paragraph(b, style="List Bullet")

    # Methodology
    doc.add_heading("Proposed Methodology", level=1)
    doc.add_paragraph(
        "Our approach follows a proven three-phase methodology that ensures "
        "systematic delivery with measurable checkpoints at each stage."
    )

    doc.add_heading("Phase 1: Discovery & Assessment", level=2)
    doc.add_paragraph(
        "Comprehensive audit of existing systems, identification of gaps, "
        "and stakeholder alignment workshops to define success criteria."
    )

    doc.add_heading("Phase 2: Implementation", level=2)
    doc.add_paragraph(
        "Agile delivery of core platform capabilities with iterative testing "
        "and continuous integration/deployment practices."
    )

    doc.add_heading("Phase 3: Optimization & Scale", level=2)
    doc.add_paragraph(
        "Performance tuning, user training, and strategic expansion of "
        "the solution across additional business units."
    )

    # Table
    doc.add_heading("Timeline & Investment", level=1)
    table = doc.add_table(rows=4, cols=3, style="Table Grid")
    headers = ["Phase", "Duration", "Investment"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    data = [
        ["Discovery", "4 weeks", "$50,000"],
        ["Implementation", "12 weeks", "$200,000"],
        ["Optimization", "6 weeks", "$75,000"],
    ]
    for r, row_data in enumerate(data, 1):
        for c, val in enumerate(row_data):
            table.rows[r].cells[c].text = val

    # Conclusion
    doc.add_heading("Conclusion", level=1)
    doc.add_paragraph(
        "We are confident that this proposal represents the most effective path "
        "forward for achieving your organization's technology transformation goals. "
        "Our team is prepared to begin immediately upon approval."
    )

    doc.add_paragraph("")
    doc.add_paragraph("For questions or clarifications, please contact:")
    doc.add_paragraph("John Smith, VP of Technology — jsmith@acme.com")

    # Save
    out = Path(__file__).resolve().parent.parent / "sample_templates"
    out.mkdir(exist_ok=True)
    path = out / "Company_Proposal.docx"
    doc.save(str(path))
    print(f"[+] Created: {path}")
    return str(path)


def create_sample_pptx():
    """Create sample Company_Template.pptx template."""
    prs = Presentation()
    prs.slide_width = PptInches(13.333)
    prs.slide_height = PptInches(7.5)

    # ── Slide 1: Title ──
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    for ph in slide.placeholders:
        if ph.placeholder_format.type in (1, 13):
            ph.text = "Technology Innovation Report"
            for run in ph.text_frame.paragraphs[0].runs:
                run.font.name = "Calibri"
                run.font.size = PptPt(40)
                run.font.bold = True
                run.font.color.rgb = PptRGBColor(0x1A, 0x56, 0xDB)
        elif ph.placeholder_format.type in (2, 15):
            ph.text = "Acme Corporation — Q3 2024"
            for run in ph.text_frame.paragraphs[0].runs:
                run.font.name = "Calibri"
                run.font.size = PptPt(20)
                run.font.color.rgb = PptRGBColor(0x4B, 0x55, 0x63)

    # ── Slides 2-6: Content ──
    content_slides = [
        {
            "title": "Agenda",
            "bullets": [
                "Market Overview & Trends",
                "Technology Assessment",
                "Proposed Solution Architecture",
                "Implementation Roadmap",
                "Investment & ROI Analysis",
            ],
        },
        {
            "title": "Market Overview",
            "bullets": [
                "Global digital transformation market: $3.4T by 2026",
                "AI adoption accelerating across all industries",
                "Cloud-native becoming the default architecture",
                "Cybersecurity spend increasing 15% YoY",
            ],
        },
        {
            "title": "Key Technology Trends",
            "bullets": [
                "Generative AI reshaping content and code creation",
                "Edge computing enabling real-time processing",
                "Zero-trust security frameworks gaining traction",
                "Low-code platforms democratizing development",
            ],
        },
        {
            "title": "Proposed Architecture",
            "bullets": [
                "Microservices-based platform on Kubernetes",
                "AI/ML pipeline for predictive analytics",
                "Real-time data streaming with event-driven design",
                "Multi-cloud deployment for resilience",
                "API-first approach for ecosystem integration",
            ],
        },
        {
            "title": "Next Steps",
            "bullets": [
                "Executive alignment workshop — Week 1",
                "Technical deep-dive sessions — Week 2-3",
                "Proof of concept development — Week 4-8",
                "Pilot deployment — Week 9-12",
                "Full-scale rollout planning — Week 13+",
            ],
        },
    ]

    for cs in content_slides:
        layout = prs.slide_layouts[1]  # Title and Content
        slide = prs.slides.add_slide(layout)
        for ph in slide.placeholders:
            if ph.placeholder_format.type in (1, 13):
                ph.text = cs["title"]
                for run in ph.text_frame.paragraphs[0].runs:
                    run.font.name = "Calibri"
                    run.font.size = PptPt(28)
                    run.font.bold = True
                    run.font.color.rgb = PptRGBColor(0x1A, 0x56, 0xDB)
            elif ph.placeholder_format.type == 2:
                tf = ph.text_frame
                tf.clear()
                for i, bullet in enumerate(cs["bullets"]):
                    para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    para.text = bullet
                    para.level = 0
                    for run in para.runs:
                        run.font.name = "Calibri"
                        run.font.size = PptPt(18)
                        run.font.color.rgb = PptRGBColor(0x2C, 0x3E, 0x50)

    # ── Slide 7: Thank You ──
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    for ph in slide.placeholders:
        if ph.placeholder_format.type in (1, 13):
            ph.text = "Thank You"
            for run in ph.text_frame.paragraphs[0].runs:
                run.font.name = "Calibri"
                run.font.size = PptPt(40)
                run.font.bold = True
                run.font.color.rgb = PptRGBColor(0x1A, 0x56, 0xDB)
        elif ph.placeholder_format.type in (2, 15):
            ph.text = "Questions & Discussion"

    out = Path(__file__).resolve().parent.parent / "sample_templates"
    out.mkdir(exist_ok=True)
    path = out / "Company_Template.pptx"
    prs.save(str(path))
    print(f"[+] Created: {path}")
    return str(path)


if __name__ == "__main__":
    print("Generating sample templates...\n")
    create_sample_docx()
    create_sample_pptx()
    print("\n[SUCCESS] Done! Sample templates are in: sample_templates/")
